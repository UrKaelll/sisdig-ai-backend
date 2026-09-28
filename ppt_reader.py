from pptx import Presentation


def read_ppt(file_path):
    presentation = Presentation(file_path)

    slides_data = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    texts.append(text)

        slides_data.append({
            "slide": slide_number,
            "content": texts
        })

    return slides_data